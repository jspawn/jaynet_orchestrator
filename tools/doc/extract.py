"""doc.extract — pull text out of PDF / Excel / DOCX files in one call.

The light lane, deliberately: pypdf reads PDFs that have a text layer,
openpyxl reads .xlsx, and .docx is a zip of XML the stdlib can strip. Both
third-party deps are pure Python and optional (requirements-tools.txt),
imported lazily — a missing one is a clear error, never a crash.

What this is NOT: OCR. A scanned PDF has no text layer and yields almost
nothing here — the tool says so and points at the `pdf` skill (the
pypdfium2/rapidocr throwaway-venv path). Layout-perfect tables aren't the
goal either; rows come out pipe-joined. If a document's structure IS the
question, that is the docling-class plugin's job (parked in
ToDos_for_later.md), not this tool's.
"""

from __future__ import annotations

import re
from pathlib import Path
from xml.sax.saxutils import unescape

from runtime.tool_base import Tool, ToolContext, ToolResult, resolve_in_roots, work_roots

CONVERTIBLE = (".pdf", ".xlsx", ".docx")
_MAX_CHARS = 200_000     # a bounded excerpt; the file stays for deeper slicing


def extract_text(p: Path) -> tuple[str, dict]:
    """(text, meta) for a convertible document. Raises ValueError with a
    clear, model-actionable message on missing deps / unreadable files."""
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _pdf_text(p)
    if suffix == ".xlsx":
        return _xlsx_text(p)
    if suffix == ".docx":
        return _docx_text(p)
    raise ValueError(f"unsupported type {suffix!r} — convertible: "
                     + ", ".join(CONVERTIBLE))


def _pdf_text(p: Path) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ValueError("pypdf not installed — add it with `uv pip install "
                         "--python .venv/bin/python -r requirements-tools.txt`")
    try:
        reader = PdfReader(str(p))
        pages = [(pg.extract_text() or "").strip() for pg in reader.pages]
    except Exception as e:
        raise ValueError(f"PDF unreadable: {type(e).__name__}: {e}")
    meta = {"pages": len(pages)}
    text = "\n\n".join(f"── page {i + 1} ──\n{t}" for i, t in enumerate(pages))
    if len(pages) and sum(len(t) for t in pages) / len(pages) < 40:
        meta["warning"] = ("almost no text layer — likely a scanned PDF. "
                           "Load the `pdf` skill for the OCR path instead.")
    return text, meta


def _xlsx_text(p: Path) -> tuple[str, dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("openpyxl not installed — add it with `uv pip install "
                         "--python .venv/bin/python -r requirements-tools.txt`")
    try:
        wb = load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = [" | ".join("" if c is None else str(c) for c in row)
                    for row in ws.iter_rows(values_only=True)]
            parts.append(f"== sheet: {ws.title} ==\n" + "\n".join(rows))
        wb.close()
    except Exception as e:
        raise ValueError(f"xlsx unreadable: {type(e).__name__}: {e}")
    return "\n\n".join(parts), {"sheets": len(parts)}


def _docx_text(p: Path) -> tuple[str, dict]:
    import zipfile
    try:
        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except Exception as e:
        raise ValueError(f"docx unreadable: {type(e).__name__}: {e}")
    # Paragraph and break tags become newlines, everything else drops.
    xml = re.sub(r"</w:p>|<w:br\s*/>", "\n", xml)
    text = unescape(re.sub(r"<[^>]+>", "", xml))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text, {"paragraphs": text.count("\n") + 1 if text else 0}


class DocExtract(Tool):
    name = "doc.extract"
    description = ("Extract text from a document file (.pdf with a text layer, "
                   ".xlsx, .docx) in one call — no venv, no job. Returns bounded "
                   "text plus page/sheet counts. For huge results, slice the file "
                   "with code.run instead of re-extracting. Scanned PDFs (no text "
                   "layer) are flagged — load the `pdf` skill for OCR. rag.index "
                   "converts these formats automatically when given a path.")
    private = True
    read_only = True
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string",
                     "description": "Document file inside the workspace roots "
                                    "(.pdf / .xlsx / .docx)."},
        },
        "required": ["path"],
    }

    async def execute(self, args: dict, ctx: ToolContext) -> ToolResult:
        try:
            p = resolve_in_roots(work_roots(ctx), args["path"])
        except (PermissionError, FileNotFoundError) as e:
            return ToolResult(status="error", result=None, error=str(e))
        try:
            text, meta = extract_text(p)
        except ValueError as e:
            return ToolResult(status="error", result=None, error=str(e))
        truncated = len(text) > _MAX_CHARS
        if truncated:
            text = text[:_MAX_CHARS]
        return ToolResult(status="ok", tool_name=self.name, result={
            "path": str(p), "chars": len(text), "truncated": truncated,
            **meta,
            "text": text,
            **({"note": f"excerpt capped at {_MAX_CHARS} chars — slice the "
                        "file with code.run for the rest"} if truncated else {}),
        })
