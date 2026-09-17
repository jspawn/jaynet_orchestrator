"""doc.extract: light-lane document text extraction (pdf/xlsx/docx) +
rag.index auto-conversion of document paths."""
import asyncio
import zipfile

import pytest

from runtime.tool_base import ToolContext
from tools.doc.extract import DocExtract, extract_text

CFG = {"orchestrator": {"model": "m", "litellm_base": "http://x:4000"}}


def _ctx(root):
    return ToolContext(request_id="t", config=CFG, budget=None,
                       work_root=str(root))


# A minimal one-page PDF with a real text layer, xref table included
# (pypdf 6.x refuses header-only files).
def _make_pdf(text: str) -> bytes:
    stream = f"BT /F1 24 Tf 100 700 Td ({text}) Tj ET".encode()
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>\nstream\n"
        + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<</Size {len(objs) + 1}/Root 1 0 R>>\n"
            f"startxref\n{xref_pos}\n%%EOF\n").encode()
    return bytes(out)


def test_pdf_text_layer(tmp_path):
    pytest.importorskip("pypdf")
    f = tmp_path / "doc.pdf"
    f.write_bytes(_make_pdf("hello jaynet"))
    text, meta = extract_text(f)
    assert "hello jaynet" in text and meta["pages"] == 1


def test_xlsx_cells(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "map"
    ws.append(["cell", "value"])
    ws.append(["A1", "START"])
    wb.save(tmp_path / "m.xlsx")
    text, meta = extract_text(tmp_path / "m.xlsx")
    assert "== sheet: map ==" in text and "A1 | START" in text
    assert meta["sheets"] == 1


def test_docx_stdlib(tmp_path):
    f = tmp_path / "d.docx"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("word/document.xml",
                   "<w:body><w:p><w:r><w:t>first para</w:t></w:r></w:p>"
                   "<w:p><w:r><w:t>second &amp; done</w:t></w:r></w:p>"
                   "</w:body>")
    text, meta = extract_text(f)
    assert "first para" in text and "second & done" in text


def test_extract_unknown_suffix(tmp_path):
    f = tmp_path / "x.bin"
    f.write_bytes(b"\x00")
    with pytest.raises(ValueError, match="unsupported type"):
        extract_text(f)


def test_tool_confined_to_roots(tmp_path):
    res = asyncio.run(DocExtract().execute({"path": "/etc/passwd"}, _ctx(tmp_path)))
    assert res.status == "error"


def test_tool_ok(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["k", "v"])
    wb.save(tmp_path / "t.xlsx")
    res = asyncio.run(DocExtract().execute({"path": "t.xlsx"}, _ctx(tmp_path)))
    assert res.status == "ok" and "k | v" in res.result["text"]


# ---- rag.index auto-conversion ----

def test_rag_index_converts_xlsx(ctx, project, tmp_path, monkeypatch):
    """A document path indexes as extracted TEXT — read_text on an xlsx would
    chunk binary garbage into the store."""
    pytest.importorskip("openpyxl")
    import openpyxl
    from conftest import run

    from tools.rag import store as rag_store

    wb = openpyxl.Workbook()
    wb.active.append(["name", "role"])
    wb.active.append(["jaynet", "orchestrator"])
    wb.save(project / "team.xlsx")

    seen = []

    async def fake_embed(texts, _ctx):
        seen.extend(texts)
        return [[1.0, 0.0] for _ in texts]

    monkeypatch.setattr(rag_store, "_embed", fake_embed)
    cfg = {"tools": {"rag": {"db_path": str(tmp_path / "rag.db")},
                     "fs": {"allowed_roots": [str(project)]}}}
    c = ctx(config=cfg, work_root=str(project))
    r = run(rag_store.RagIndex().execute(
        {"collection": "c", "path": "team.xlsx"}, c))
    assert r.status == "ok" and r.result["chunks_indexed"] >= 1
    assert any("orchestrator" in t for t in seen)
